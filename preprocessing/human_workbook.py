"""Export the human anchor worksheet as an Excel workbook for annotation.

The CSV is the source of truth, but it opens badly in Excel on a French-locale
Windows machine: no UTF-8 BOM, so accents render as mojibake, and a comma
separator where Excel expects a semicolon, so every field lands in column A. The
workbook avoids both, and adds what makes hand-labelling reliable: a dropdown
restricted to the five labels, the rubric on a second sheet, and a progress count.

It contains headlines only. No model label and no v1 label is written into it --
the anchor is worthless if the annotator can see what the models said.

    python preprocessing/human_workbook.py            # CSV -> XLSX
    (label in Excel, save)
    python preprocessing/merge_human.py --worksheet data/curated/sentiment_gold_v2_human.xlsx
"""

import argparse
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from gold import LABELS

ROOT = Path(__file__).resolve().parent.parent
CURATED = ROOT / "data" / "curated"
DEFAULT_INPUT = CURATED / "sentiment_gold_v2_human.csv"
DEFAULT_OUTPUT = CURATED / "sentiment_gold_v2_human.xlsx"
SHEET = "annotate"

COLUMNS = ["gold_item_id", "source", "published_date", "headline_clean",
           "human_label", "human_notes"]
WIDTHS = {"gold_item_id": 12, "source": 18, "published_date": 12,
          "headline_clean": 90, "human_label": 16, "human_notes": 45}

RUBRIC = [
    ("THE QUESTION", ""),
    ("", "Would a Tunisian equity investor, reading only this headline, become more "
         "or less optimistic about the near-term value of Tunisian listed companies?"),
    ("", "Judge expected MARKET IMPACT only - not whether the news is pleasant, "
         "economic-sounding, or upbeat."),
    ("", ""),
    ("LABELS", ""),
    ("very_negative", "Large, direct hit to earnings, solvency or cost of capital: sovereign "
                      "downgrade, banking crisis, default, major devaluation, sharp index fall."),
    ("negative", "Clear adverse effect of ordinary size: rising inflation, rate hike, falling "
                 "profits, widening deficit, strike at a listed firm."),
    ("neutral", "No clear directional implication, or effects offset: unchanged policy, "
                "procedural and announcement news, off-topic items. THE DEFAULT."),
    ("positive", "Clear favourable effect of ordinary size: falling inflation, rate cut, "
                 "rising profits, new financing secured, improved outlook."),
    ("very_positive", "Large, direct improvement: major investment inflow, sovereign upgrade, "
                      "debt relief, sharp index rise, record results at a large listed company."),
    ("", ""),
    ("RULES", ""),
    ("", "Neutral is the default. Choose a direction only if you can say which way prices "
         "move and why. Reserve the very_ labels for large or systemic items."),
    ("", "Foreign news is neutral unless it transmits to Tunisia (oil, EU demand, Fed/ECB "
         "rates, remittances, tourism, grain). Another country's domestic affairs: neutral."),
    ("", ""),
    ("TRAPS", ""),
    ("", "'hausse' / 'increases' is not automatically positive: rising inflation, debt, "
         "unemployment or money supply are negative or neutral."),
    ("", "A draft law, plan or project under preparation has not happened yet: neutral."),
    ("", "Topic is not sentiment: a headline that merely concerns the economy is not positive."),
    ("", "'maintient' / 'inchangé' means no change: neutral."),
    ("", "A question or a wish ('Vivement la baisse du taux') states no fact: neutral."),
    ("", "A headline reporting the index's own move takes the direction it reports "
         "('grignote', 's'effrite' are small but still directional)."),
]


def build_workbook(input_path: Path = DEFAULT_INPUT,
                   output_path: Path = DEFAULT_OUTPUT) -> Path:
    frame = pd.read_csv(input_path, keep_default_na=False, dtype=str)
    missing = set(COLUMNS) - set(frame.columns)
    if missing:
        raise ValueError(f"Worksheet is missing columns: {sorted(missing)}")

    book = Workbook()
    sheet = book.active
    sheet.title = SHEET
    sheet.append(COLUMNS)
    for row in frame[COLUMNS].itertuples(index=False):
        sheet.append(list(row))

    header_fill = PatternFill("solid", fgColor="DDDDDD")
    label_fill = PatternFill("solid", fgColor="FFF4CC")
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
    for letter, column in zip("ABCDEF", COLUMNS):
        sheet.column_dimensions[letter].width = WIDTHS[column]
    last = len(frame) + 1
    for row in sheet.iter_rows(min_row=2, max_row=last):
        row[3].alignment = Alignment(wrap_text=True, vertical="top")
        row[4].fill = label_fill
        row[5].alignment = Alignment(wrap_text=True, vertical="top")
    sheet.freeze_panes = "E2"

    dropdown = DataValidation(type="list", formula1='"' + ",".join(LABELS) + '"',
                              allow_blank=True, showDropDown=False)
    dropdown.error = f"Choose one of: {', '.join(LABELS)}"
    dropdown.errorTitle = "Invalid label"
    dropdown.showErrorMessage = True
    sheet.add_data_validation(dropdown)
    dropdown.add(f"E2:E{last}")

    sheet["H1"] = "progress"
    sheet["H1"].font = Font(bold=True)
    sheet["H2"] = f'=COUNTA(E2:E{last})&" / {len(frame)}"'
    sheet.column_dimensions["H"].width = 12

    rubric = book.create_sheet("rubric")
    for key, text in RUBRIC:
        rubric.append([key, text])
    rubric.column_dimensions["A"].width = 16
    rubric.column_dimensions["B"].width = 110
    for row in rubric.iter_rows():
        row[0].font = Font(bold=True)
        row[1].alignment = Alignment(wrap_text=True, vertical="top")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    book.save(output_path)
    return output_path


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__,
                                     formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    path = build_workbook(args.input, args.output)
    print(f"Wrote {path}")


if __name__ == "__main__":
    main()
