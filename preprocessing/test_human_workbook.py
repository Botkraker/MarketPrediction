import pandas as pd
from openpyxl import load_workbook

from gold import LABELS
from human_workbook import build_workbook
from merge_human import merge_human


def _worksheet_csv(tmp_path):
    path = tmp_path / "human.csv"
    pd.DataFrame({
        "gold_item_id": ["gold-00001", "gold-00002"],
        "source": ["ilboursa", "kapitalis"],
        "published_date": ["2020-01-02", "2021-03-04"],
        "headline_clean": ["L’inflation recule à 5,2 %", "Le Tunindex s’effrite"],
        "human_label": ["", ""],
        "human_notes": ["", ""],
    }).to_csv(path, index=False)
    return path


def test_workbook_has_headlines_dropdown_and_no_model_labels(tmp_path):
    xlsx = build_workbook(_worksheet_csv(tmp_path), tmp_path / "human.xlsx")
    book = load_workbook(xlsx)
    sheet = book["annotate"]

    header = [cell.value for cell in sheet[1]][:6]
    assert header == ["gold_item_id", "source", "published_date", "headline_clean",
                      "human_label", "human_notes"]
    assert sheet["D2"].value == "L’inflation recule à 5,2 %"   # accents survive
    validation = sheet.data_validations.dataValidation[0]
    assert validation.formula1 == '"' + ",".join(LABELS) + '"'
    assert "E2:E3" in str(validation.sqref)
    assert "rubric" in book.sheetnames
    flat = " ".join(str(c.value) for row in sheet.iter_rows() for c in row if c.value)
    assert "annotator_" not in flat and "v1_label" not in flat


def test_labelled_workbook_merges(tmp_path):
    xlsx = build_workbook(_worksheet_csv(tmp_path), tmp_path / "human.xlsx")
    book = load_workbook(xlsx)
    book["annotate"]["E2"] = "positive"
    book["annotate"]["F2"] = "disinflation"
    book.save(xlsx)

    target = tmp_path / "gold.csv"
    pd.DataFrame({"gold_item_id": ["gold-00001", "gold-00002"],
                  "annotator_1_label": ["neutral", "negative"]}).to_csv(target, index=False)
    merge_human(xlsx, target, tmp_path / "meta.json")

    out = pd.read_csv(target, keep_default_na=False)
    assert out["annotator_3_label"].tolist() == ["positive", ""]
    assert out["annotator_3_notes"].tolist() == ["disinflation", ""]
